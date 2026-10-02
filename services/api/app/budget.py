"""An account's monthly labor plan against its actuals (migration 046): the account Budget tab.

The plan is entered in Admin > Budgets (pasted from the account's budget workbook): site labor,
overhead labor, revenue and supplies per month, with the calendar it was built on kept as details.

Actuals use the same split as the plan:
* site labor: the account's site jobs; overhead: its catch-all jobs (e.g. Plano's job 800);
* events: its non-billed jobs (e.g. Plano's job 896, Special Events), shown separately and left out
  of the actual total, as budgets are set without them;
* a month with job cost labor for the account's sites reads job cost (basis job_cost); a month
  without, up to the current one, reads the month rollup from timekeeping (app/month.py; basis
  pay_report, or estimate when any of it is the trailing-rate estimate). Later months have no actual.

A month still running (its last day not yet past) carries its actual to date, flagged in_progress,
and no variance: it cannot be compared with a whole month's plan.

Labor % is labor over revenue: the planned revenue for the plan, the actual billing (job cost, else
the plan's revenue) for the actual.
"""
from __future__ import annotations

from datetime import date, timedelta
from decimal import Decimal
from typing import Any

FIELDS = ("site_labor", "overhead_labor", "revenue", "supplies")
DETAIL_KEYS = ("school_days", "staff_days", "closure_days", "summer_days", "stat_holidays")


MAX_SHEET_ROWS = 2000
MAX_SHEET_COLUMNS = 60


def _cell(value: Any) -> str:
    """One cell as the paste parser reads it: dates ISO, whole numbers without .0, text quoted when it holds a
    tab, line break or quote (Excel's own copy format)."""
    from datetime import datetime as dt

    if value is None:
        return ""
    if isinstance(value, dt):
        value = value.date()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = str(value)
    if any(c in text for c in "\t\n\r\""):
        return '"' + text.replace('"', '""') + '"'
    return text


def workbook_sheets(content: bytes) -> list[dict[str, str]]:
    """Each sheet of an Excel workbook as tab-separated text (values, not formulas), for the budget parser in
    the browser, which finds the monthly plan and the weekly calendar by their headers. Empty sheets are left out."""
    import io

    from openpyxl import load_workbook

    book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    out = []
    try:
        for sheet in book.worksheets:
            lines = []
            for row in sheet.iter_rows(max_row=MAX_SHEET_ROWS, max_col=MAX_SHEET_COLUMNS, values_only=True):
                cells = [_cell(v) for v in row]
                while cells and cells[-1] == "":
                    cells.pop()
                if cells:
                    lines.append("\t".join(cells))
            if lines:
                out.append({"name": sheet.title, "text": "\n".join(lines)})
    finally:
        book.close()
    return out


def _f(value: Any) -> float | None:
    return None if value is None else float(value)


def validate(rows: Any) -> list[dict[str, Any]]:
    """Plan rows as saved: month YYYY-MM, non-negative amounts, numeric details. ValueError otherwise."""
    if not isinstance(rows, list) or not rows:
        raise ValueError("months must be a non-empty list")
    out, seen = [], set()
    for n, row in enumerate(rows, start=1):
        if not isinstance(row, dict):
            raise ValueError(f"row {n} must be an object")
        try:
            month = date.fromisoformat(f"{str(row.get('month', ''))[:7]}-01")
        except ValueError:
            raise ValueError(f"row {n}: month must be YYYY-MM") from None
        if month in seen:
            raise ValueError(f"row {n}: {month:%Y-%m} appears twice")
        seen.add(month)
        clean: dict[str, Any] = {"month": month}
        for key in FIELDS:
            value = row.get(key)
            if value is not None and (isinstance(value, bool) or not isinstance(value, (int, float)) or value < 0):
                raise ValueError(f"row {n}: {key} must be a non-negative number")
            clean[key] = None if value is None else Decimal(str(round(value, 2)))
        if clean["site_labor"] is None and clean["overhead_labor"] is None:
            raise ValueError(f"row {n}: site or overhead labor is required")
        details = row.get("details") or {}
        if not isinstance(details, dict) or any(k not in DETAIL_KEYS or isinstance(v, bool) or not isinstance(v, (int, float)) for k, v in details.items()):
            raise ValueError(f"row {n}: details takes numbers for {', '.join(DETAIL_KEYS)}")
        clean["details"] = details
        out.append(clean)
    return out


WEEK_FIELDS = ("site_labor", "overhead_labor", "holiday_labor")
WEEK_DETAIL_KEYS = ("school_days", "staff_days", "closure_days", "summer_days", "stat_holidays")


def validate_weeks(rows: Any) -> list[dict[str, Any]]:
    """Weekly plan rows: week_end a Sunday (YYYY-MM-DD), non-negative amounts. ValueError otherwise."""
    if not isinstance(rows, list):
        raise ValueError("weeks must be a list")
    out, seen = [], set()
    for n, row in enumerate(rows, start=1):
        if not isinstance(row, dict):
            raise ValueError(f"week {n} must be an object")
        try:
            week_end = date.fromisoformat(str(row.get("week_end", ""))[:10])
        except ValueError:
            raise ValueError(f"week {n}: week_end must be YYYY-MM-DD") from None
        if week_end.isoweekday() != 7:
            raise ValueError(f"week {n}: {week_end} is not a Sunday (weeks end on Sunday)")
        if week_end in seen:
            raise ValueError(f"week {n}: {week_end} appears twice")
        seen.add(week_end)
        clean: dict[str, Any] = {"week_end": week_end}
        for key in WEEK_FIELDS:
            value = row.get(key) or 0
            if isinstance(value, bool) or not isinstance(value, (int, float)) or value < 0:
                raise ValueError(f"week {n}: {key} must be a non-negative number")
            clean[key] = Decimal(str(round(value, 2)))
        details = row.get("details") or {}
        if not isinstance(details, dict) or any(k not in WEEK_DETAIL_KEYS or isinstance(v, bool) or not isinstance(v, (int, float)) for k, v in details.items()):
            raise ValueError(f"week {n}: details takes numbers for {', '.join(WEEK_DETAIL_KEYS)}")
        clean["details"] = details
        out.append(clean)
    return out


def save_weeks(cursor: Any, slug: str, rows: list[dict[str, Any]], actor: str) -> None:
    from psycopg.types.json import Jsonb

    for r in rows:
        cursor.execute(
            """
            INSERT INTO ops.account_budget_week (account_slug, week_end, site_labor, overhead_labor, holiday_labor, details, updated_by)
            VALUES (%s, %s, %s, %s, %s, %s, %s)
            ON CONFLICT (account_slug, week_end) DO UPDATE SET site_labor = EXCLUDED.site_labor, overhead_labor = EXCLUDED.overhead_labor,
              holiday_labor = EXCLUDED.holiday_labor, details = EXCLUDED.details, updated_at = now(), updated_by = EXCLUDED.updated_by
            """,
            (slug, r["week_end"], r["site_labor"], r["overhead_labor"], r["holiday_labor"], Jsonb(r["details"]), actor),
        )


def weeks(cursor: Any, slug: str) -> list[dict[str, Any]]:
    cursor.execute("SELECT week_end, site_labor, overhead_labor, holiday_labor, details FROM ops.account_budget_week "
                   "WHERE account_slug = %s ORDER BY week_end", (slug,))
    return [{"week_end": r["week_end"].isoformat(), "site": float(r["site_labor"]), "overhead": float(r["overhead_labor"]),
             "holiday": float(r["holiday_labor"]), "details": r["details"] or {}} for r in cursor.fetchall()]


def save(cursor: Any, slug: str, rows: list[dict[str, Any]], actor: str) -> None:
    from psycopg.types.json import Jsonb

    for r in rows:
        cursor.execute(
            """
            INSERT INTO ops.account_budget_month (account_slug, month, site_labor, overhead_labor, revenue, supplies, details, updated_by)
            VALUES (%s, %s, %s, %s, %s, %s, %s, %s)
            ON CONFLICT (account_slug, month) DO UPDATE SET site_labor = EXCLUDED.site_labor, overhead_labor = EXCLUDED.overhead_labor,
              revenue = EXCLUDED.revenue, supplies = EXCLUDED.supplies, details = EXCLUDED.details, updated_at = now(), updated_by = EXCLUDED.updated_by
            """,
            (slug, r["month"], r["site_labor"], r["overhead_labor"], r["revenue"], r["supplies"], Jsonb(r["details"]), actor),
        )


def plan(cursor: Any, slug: str) -> list[dict[str, Any]]:
    cursor.execute("SELECT month, site_labor, overhead_labor, revenue, supplies, details, updated_at, updated_by FROM ops.account_budget_month "
                   "WHERE account_slug = %s ORDER BY month", (slug,))
    return [dict(r) for r in cursor.fetchall()]


def job_cost_actuals(cursor: Any, slug: str, months: list[date]) -> dict[date, dict[str, float]]:
    """Job cost labor and revenue per month by job role, for months whose site jobs carry job cost labor."""
    cursor.execute(
        """
        SELECT jc.month, aj.role, sum(jc.direct_labor) AS labor, sum(jc.revenue) AS revenue
        FROM mart.v_job_cost_month_effective jc
        JOIN ops.account_job aj ON aj.company = jc.company AND aj.job_number = jc.job_number
        WHERE aj.account_slug = %s AND jc.month = ANY(%s)
        GROUP BY 1, 2
        """,
        (slug, months),
    )
    out: dict[date, dict[str, float]] = {}
    for r in cursor.fetchall():
        m = out.setdefault(r["month"], {"site": 0.0, "overhead": 0.0, "events": 0.0, "revenue": 0.0})
        key = {"site": "site", "catch_all": "overhead", "non_billed": "events"}.get(r["role"], "site")
        m[key] += float(r["labor"] or 0)
        m["revenue"] += float(r["revenue"] or 0)
    return {k: v for k, v in out.items() if v["site"] > 0}


def rollup_actual(rows: list[dict[str, Any]], slug: str) -> dict[str, Any]:
    """One month's actual from month rollup rows (app/month.py) of every account."""
    mine = [r for r in rows if r.get("account_slug") == slug]
    sums = {"site": 0.0, "overhead": 0.0, "events": 0.0, "revenue": 0.0}
    estimated = False
    for r in mine:
        key = {"site": "site", "catch_all": "overhead", "non_billed": "events"}.get(r["role"], "site")
        sums[key] += float(r["labor"] or 0)
        sums["revenue"] += float(r.get("revenue_month_amount") or 0)
        estimated = estimated or (r["labor_basis"] != "pay_report" and float(r["labor"] or 0) > 0)
    return {**sums, "basis": "estimate" if estimated else "pay_report"}


def report(cursor: Any, slug: str, today: date | None = None, rollup: Any = None) -> list[dict[str, Any]]:
    """Every planned month with its actual (when the month has started), variance and labor %."""
    today = today or date.today()
    months = plan(cursor, slug)
    started = [m["month"] for m in months if m["month"] <= today]
    closed = job_cost_actuals(cursor, slug, started) if started else {}
    out = []
    for m in months:
        site, overhead = _f(m["site_labor"]) or 0.0, _f(m["overhead_labor"]) or 0.0
        budget_total = site + overhead
        revenue = _f(m["revenue"])
        actual = None
        if m["month"] in closed:
            actual = {**closed[m["month"]], "basis": "job_cost"}
        elif m["month"] <= today and rollup is not None:
            actual = rollup_actual(rollup(m["month"]), slug)
        last = (m["month"].replace(day=28) + timedelta(days=4)).replace(day=1) - timedelta(days=1)
        in_progress = m["month"] <= today <= last
        row: dict[str, Any] = {
            "month": m["month"].isoformat(), "in_progress": in_progress, "details": m["details"] or {}, "supplies": _f(m["supplies"]),
            "budget": {"site": site, "overhead": overhead, "total": budget_total, "revenue": revenue,
                       "labor_pct": budget_total / revenue if revenue else None},
            "actual": None, "variance": None,
        }
        if actual is not None:
            total = actual["site"] + actual["overhead"]
            billing = actual["revenue"] or revenue or 0.0
            row["actual"] = {"site": round(actual["site"], 2), "overhead": round(actual["overhead"], 2), "events": round(actual["events"], 2),
                             "total": round(total, 2), "revenue": round(billing, 2), "basis": actual["basis"],
                             "labor_pct": total / billing if billing else None}
            row["variance"] = None if in_progress else {"total": round(total - budget_total, 2), "site": round(actual["site"] - site, 2),
                               "overhead": round(actual["overhead"] - overhead, 2),
                               "pct": (total - budget_total) / budget_total if budget_total else None,
                               "points": (row["actual"]["labor_pct"] - row["budget"]["labor_pct"])
                               if row["actual"]["labor_pct"] is not None and row["budget"]["labor_pct"] is not None else None}
        out.append(row)
    return out
