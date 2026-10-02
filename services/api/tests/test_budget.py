"""An account's monthly labor plan against actuals (app/budget.py), with Plano ISD's 2026 numbers."""
from __future__ import annotations

from datetime import date, datetime, timezone
from decimal import Decimal

import pytest

from app import budget

AUG = {"month": date(2026, 8, 1), "site_labor": Decimal("653907"), "overhead_labor": Decimal("40230"), "revenue": Decimal("1026956"),
       "supplies": Decimal("125000"), "details": {"school_days": 15, "summer_days": 6}, "updated_at": datetime.now(timezone.utc), "updated_by": "x"}
SEP = {**AUG, "month": date(2026, 9, 1), "site_labor": Decimal("817385"), "overhead_labor": Decimal("42146"), "details": {"school_days": 21}}
OCT = {**AUG, "month": date(2026, 10, 1), "site_labor": Decimal("739538"), "overhead_labor": Decimal("42146")}


def test_validate_reads_the_pasted_plan():
    rows = budget.validate([{"month": "2026-08", "site_labor": 653907, "overhead_labor": 40230, "revenue": 1026956, "supplies": 125000,
                             "details": {"school_days": 15, "summer_days": 6}}])
    assert rows[0]["month"] == date(2026, 8, 1) and rows[0]["site_labor"] == Decimal("653907")
    for bad in ([], [{"month": "Aug"}], [{"month": "2026-08"}], [{"month": "2026-08", "site_labor": -1}],
                [{"month": "2026-08", "site_labor": 1}, {"month": "2026-08-15", "site_labor": 2}],
                [{"month": "2026-08", "site_labor": 1, "details": {"rain_days": 2}}]):
        with pytest.raises(ValueError):
            budget.validate(bad)


def test_august_matches_the_email(monkeypatch):
    """Sites $772,172 + job 800 $133,693 = $905,865 ex events (email: $905,864, 88.2%); events job 896 left out."""
    monkeypatch.setattr(budget, "plan", lambda cursor, slug: [AUG, SEP, OCT])
    monkeypatch.setattr(budget, "job_cost_actuals", lambda cursor, slug, months: {
        date(2026, 8, 1): {"site": 772172.0, "overhead": 133693.0, "events": 5199.0, "revenue": 1026956.0}})
    rollup_rows = [
        {"account_slug": "plano-isd", "role": "site", "labor": 846641.0, "labor_basis": "trailing_rate_estimate", "revenue_month_amount": 1026956.0},
        {"account_slug": "plano-isd", "role": "catch_all", "labor": 83096.0, "labor_basis": "trailing_rate_estimate", "revenue_month_amount": 0},
        {"account_slug": "plano-isd", "role": "non_billed", "labor": 33909.0, "labor_basis": "trailing_rate_estimate", "revenue_month_amount": 0},
        {"account_slug": "fedex", "role": "site", "labor": 1e6, "labor_basis": "pay_report", "revenue_month_amount": 9e6},
    ]
    asked = []
    out = budget.report(object(), "plano-isd", today=date(2026, 10, 2), rollup=lambda m: asked.append(m) or rollup_rows)
    aug, sep, octo = out
    assert aug["budget"]["total"] == 694137 and round(aug["budget"]["labor_pct"], 3) == 0.676
    assert aug["actual"]["total"] == 905865 and aug["actual"]["events"] == 5199 and aug["actual"]["basis"] == "job_cost"
    assert round(aug["actual"]["labor_pct"], 3) == 0.882 and aug["variance"]["total"] == 211728
    assert (aug["variance"]["site"], aug["variance"]["overhead"]) == (118265, 93463)
    assert round(aug["variance"]["points"] * 100, 1) == 20.6
    assert sep["actual"]["basis"] == "estimate" and sep["actual"]["total"] == 929737 and sep["budget"]["total"] == 859531
    assert asked == [date(2026, 9, 1), date(2026, 10, 1)]  # job cost for August; Oct has started
    assert octo["actual"]["total"] == 929737 and octo["in_progress"] and octo["variance"] is None  # Oct 2: to date only
    assert not sep["in_progress"] and sep["variance"]["total"] == 929737 - 859531
    later = budget.report(object(), "plano-isd", today=date(2026, 9, 15), rollup=lambda m: rollup_rows)
    assert later[2]["actual"] is None and later[2]["variance"] is None  # October not started


def test_weekly_calendar_rows():
    """Plano's week ending Sep 13: four school days and Labor Day (stat holiday pay kept apart)."""
    rows = budget.validate_weeks([{"week_end": "2026-09-13", "site_labor": 155691.58, "overhead_labor": 9578.54, "holiday_labor": 37069.42,
                                   "details": {"school_days": 4, "stat_holidays": 1}}])
    assert rows[0]["week_end"] == date(2026, 9, 13) and rows[0]["holiday_labor"] == Decimal("37069.42")
    assert budget.validate_weeks([]) == []
    for bad in ([{"week_end": "2026-09-12", "site_labor": 1}], [{"week_end": "Sep 13", "site_labor": 1}],
                [{"week_end": "2026-09-13", "site_labor": -5}], [{"week_end": "2026-09-13"}, {"week_end": "2026-09-13"}]):
        with pytest.raises(ValueError):
            budget.validate_weeks(bad)


def test_a_workbook_reads_as_one_table_per_sheet():
    import io
    from datetime import datetime as dt

    from openpyxl import Workbook

    book = Workbook()
    plan = book.active
    plan.title = "FY27 plan"
    plan.append(["Month", "School\ndays", "Site labor", "Overhead labor", "Revenue"])
    plan.append(["Aug 2026", 15, 653907, 40230.0, 1026956])
    plan.append([dt(2026, 9, 1), 21, 817385.5, 42146, 1026956])
    weeks = book.create_sheet("Weekly")
    weeks.append(["Week ending", "Site labor", "Overhead labor"])
    weeks.append([dt(2026, 9, 13), 155691.58, 9578.54])
    book.create_sheet("Notes")
    buf = io.BytesIO()
    book.save(buf)
    sheets = budget.workbook_sheets(buf.getvalue())
    assert [s["name"] for s in sheets] == ["FY27 plan", "Weekly"]
    assert sheets[0]["text"].splitlines()[0] == 'Month\t"School'  # the wrapped header is quoted, as Excel copies it
    assert "Aug 2026\t15\t653907\t40230\t1026956" in sheets[0]["text"]
    assert "2026-09-01\t21\t817385.5\t42146\t1026956" in sheets[0]["text"]
    assert sheets[1]["text"].splitlines()[1] == "2026-09-13\t155691.58\t9578.54"
