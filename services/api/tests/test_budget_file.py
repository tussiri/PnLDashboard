"""Mailed labor budget workbooks (app/budget_file.py): the same reading as Admin > Budgets."""
from __future__ import annotations

import io

from openpyxl import Workbook

from app import budget_file

ACCOUNTS = [{"slug": "plano-isd", "name": "Plano ISD"}, {"slug": "fedex", "name": "FedEx"}, {"slug": "crowley-isd", "name": "Crowley ISD"}]


def test_reads_the_monthly_plan_and_the_weekly_calendar_from_a_workbook():
    book = Workbook()
    plan = book.active
    plan.title = "FY27 Labor Budget"
    plan.append(["Plano ISD FY27 labor budget"])
    plan.append(["Month", "School\ndays", "Site labor", "Overhead labor", "Total labor", "Revenue"])
    plan.append(["Jul 2026", 0, 245765.21, 41376.27, 287141.48, None])
    plan.append(["Aug 2026", 10, "$686,385.26", "41,376.27", 727761.53, 1026956])
    plan.append(["Total", None, 932150.47, 82752.54, 1014903.01, None])
    weeks = book.create_sheet("Weekly calendar")
    weeks.append(["Week ending", "Site labor", "Overhead labor", "Stat holiday labor", "School days", "Stat holidays"])
    weeks.append(["2026-09-13", 155691.58, 9578.54, 37069.42, 4, 1])
    book.create_sheet("Notes").append(["Built from the district calendar"])
    out = io.BytesIO()
    book.save(out)
    months, wk, errors = budget_file.read("Plano_ISD_FY27_labor_budget.xlsx", out.getvalue())
    assert errors == []
    assert [(m["month"], m["site_labor"], m["overhead_labor"], m["revenue"]) for m in months] == [
        ("2026-07", 245765.21, 41376.27, None), ("2026-08", 686385.26, 41376.27, 1026956.0)]
    assert months[1]["details"] == {"school_days": 10.0}
    assert wk == [{"week_end": "2026-09-13", "site_labor": 155691.58, "overhead_labor": 9578.54, "holiday_labor": 37069.42,
                   "details": {"school_days": 4.0, "stat_holidays": 1.0}}]


def test_a_total_that_disagrees_is_an_error():
    _m, _w, errors = budget_file.parse_sheet("Month\tSite labor\tOverhead labor\tTotal labor\nSep 2026\t100\t50\t200\n")
    assert errors and "does not equal total labor" in errors[0]


def test_a_sheet_without_either_table_is_not_a_budget():
    assert budget_file.parse_sheet("Invoice\tCustomer\tAmount\n1001\tACME\t500\n") == ([], [], [])
    assert not budget_file.looks_like("ar.csv", b"Invoice,Customer,Amount\n1001,ACME,500\n")


def test_the_account_comes_from_the_file_name_then_the_subject():
    assert budget_file.account_for(ACCOUNTS, "Plano_ISD_FY27_labor_budget.xlsx", None) == "plano-isd"
    assert budget_file.account_for(ACCOUNTS, "FY27_labor_budget.xlsx", "Crowley ISD budget, revised") == "crowley-isd"
    assert budget_file.account_for(ACCOUNTS, "budget.xlsx", "FY27") is None


def test_parsers_match_the_browser():
    assert [budget_file.parse_month(v) for v in ("Jul 2026", "July 2026", "2026-07", "7/2026", "07/01/2026", "Total")] == ["2026-07"] * 5 + [None]
    assert [budget_file.parse_day(v) for v in ("2026-09-13", "9/13/2026", "Sep 13, 2026", "Average")] == ["2026-09-13"] * 3 + [None]
    assert [budget_file.parse_amount(v) for v in ("$1,026,956", "(1,200)", "67.6%", "")] == [1026956.0, -1200.0, 67.6, None]


def test_an_upload_on_imports_that_is_a_budget_goes_to_the_budget_loader(monkeypatch):
    """Admin > Imports with auto-detect: a budget workbook is saved as a budget, not failed as an unknown feed."""
    import asyncio
    from contextlib import contextmanager

    from app import imports
    from app.routers import leadership

    calls: list[tuple[str, str]] = []
    monkeypatch.setattr(budget_file, "load_file", lambda conn, name, content, subject=None, origin="mail", uploaded_by=None:
                        calls.append((name, origin)) or {"kind": "budget", "status": "loaded", "import_file_id": 3, "sha256": "x"})
    monkeypatch.setattr(imports, "load_file", lambda *a, **k: (_ for _ in ()).throw(AssertionError("not a feed")))
    monkeypatch.setattr(leadership, "_actor", lambda request: "admin")

    @contextmanager
    def fake_connection():
        yield object()

    monkeypatch.setattr(leadership, "connection", fake_connection)

    class Upload:
        filename = "Plano_ISD_FY27_weekly_budget.csv"
        async def read(self, n):
            return b"Week ending,Site labor,Overhead labor\n2026-09-20,194614.48,9578.54\n"

    out = asyncio.run(leadership.upload_import(request=None, file=Upload(), kind=None, rebuild=True))
    assert calls == [("Plano_ISD_FY27_weekly_budget.csv", "upload")]
    assert out == {"file": {"kind": "budget", "status": "loaded", "import_file_id": 3}, "marts": None}
